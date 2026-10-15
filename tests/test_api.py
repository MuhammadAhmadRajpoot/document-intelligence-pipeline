import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from docintel.api.app import create_app
from tests.conftest import INVOICE_JSON
from tests.test_ingest import text_pdf
from tests.test_pipeline import PDF_TEXT


@pytest.fixture
def client(make_pipeline, settings):
    return TestClient(create_app(pipeline=make_pipeline([INVOICE_JSON]), settings=settings))


def test_health(client):
    body = client.get("/health").json()
    assert body["status"] == "ok"


def test_index_and_static_assets(client):
    assert "Document Intelligence" in client.get("/").text
    assert client.get("/static/app.js").status_code == 200


def test_extract_pdf(client):
    response = client.post("/extract", files={"file": ("inv.pdf", text_pdf(PDF_TEXT))})
    assert response.status_code == 200
    body = response.json()
    assert body["result"]["is_valid"] is True
    assert body["result"]["data"]["invoice_number"] == "INV-2025-1001"
    assert body["pages"][0]["preview"].startswith("data:image/jpeg;base64,")
    assert "pages" not in body["result"]


def test_extract_rejects_unsupported_and_empty(client):
    assert client.post("/extract", files={"file": ("a.txt", b"hello")}).status_code == 415
    assert client.post("/extract", files={"file": ("a.pdf", b"")}).status_code == 400


def test_extract_without_api_key_returns_503(settings):
    from docintel.config import Settings

    app = create_app(settings=Settings(_env_file=None, groq_api_key=None))
    response = TestClient(app).post("/extract", files={"file": ("inv.pdf", text_pdf(PDF_TEXT))})
    assert response.status_code == 503
    assert "GROQ_API_KEY" in response.json()["detail"]


def test_export_csv_and_xlsx(client):
    result = client.post("/extract", files={"file": ("inv.pdf", text_pdf(PDF_TEXT))}).json()
    payload = [result["result"]]
    csv_text = client.post("/export/csv", json=payload).text
    assert "invoice_number" in csv_text.splitlines()[0]
    xlsx = client.post("/export/xlsx", json=payload)
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames == ["documents", "line_items"]
    assert workbook["line_items"].max_row == 3  # header + 2 items
